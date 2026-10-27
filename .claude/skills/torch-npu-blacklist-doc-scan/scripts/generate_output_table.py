#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
步骤 5 辅助脚本：把"已完成文档核验"的中间表转换成本 skill 规定的三列输出表
（nodeid | 对应算子 | 不支持类型）。

本脚本只负责"从已核验数据产出格式化表格"这一步（步骤 5），不负责步骤 1-4
的筛查/追踪/核验——那些步骤需要读文档、做判断，不能脚本化，必须按
references/candidate-screening.md 和 references/trace-to-real-component.md
人工/半人工完成。

用法：
    python generate_output_table.py <输入文件路径> <输出md路径> [--sheet SHEET名]

输入文件支持 .xlsx/.xlsm（按 --sheet 指定或默认取当前激活的 sheet）或 .csv
（按扩展名自动识别，UTF-8 编码）。

列是按表头文字识别的，不是按固定位置，列序随意、允许有本脚本不认识的额外列：
    必须有：nodeid
    可选：  失败大类 / torch-api / CANN算子 / 不支持类型 / 结论 / 文档链接
中间表完全可以只有 nodeid + 不支持类型两列，或者 nodeid + 结论两列——缺哪个
可选列，就把那部分信息当成"没有"处理（会退化到别的来源，或者最终留空/标记
未识别），不是必须凑齐这 6 个可选列才能跑。表头文字必须和上面的名字逐字匹配
（大小写、首尾空格容忍），匹配不上的列名会被当成"没有这一列"而不是报错，
所以如果整张表一个可选列都没对上，大概率是列名写得不一样，看 stderr 里打印
的"识别到的列"自查。

"对应算子"解析优先级（处理"测试表面现象≠真正失败组件"的情况，
见 references/trace-to-real-component.md）：
    1. 有"结论"列，且文本开头能干净匹配"XXX文档支持类型" -> 最可信，直接用 XXX
    2. 开头匹配不上，但这一行的"文档链接"跟某条"干净匹配"的记录完全相同
       -> 复用那条记录反推出的真实算子名（同一份文档，身份应该一致）
    3. 都没有 -> 依次退回"CANN算子"列、"torch-api"列
    4. 都没有 -> 标记未识别，需人工核实

"不支持类型"规则：
    - 有"失败大类"列且其值等于 "内存格式错误"（或 --memfmt-category 自定义值）
      -> 统一写"非连续内存"
    - 否则直接取"不支持类型"列原始值（没有这一列就留空标记未填写）
"""
import argparse
import csv
import os
import re
import sys
from collections import Counter

import openpyxl

NAME_PAT = re.compile(r'^([A-Za-z_][A-Za-z0-9_.]*)文档支持类型')

REQUIRED_COLUMNS = ['nodeid']
OPTIONAL_COLUMNS = ['失败大类', 'torch-api', 'CANN算子', '不支持类型', '结论', '文档链接']
ALL_KNOWN_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


def build_column_map(header_row):
    """
    表头文字 -> 列索引，只认 ALL_KNOWN_COLUMNS 里的名字（忽略大小写/首尾空格），
    不认识的表头列直接跳过（允许中间表里有本脚本不需要的额外列，比如备注/
    置信度）。不要求 6 个可选列全部出现，也不要求出现顺序——这是本脚本唯一
    的"认列"入口，后面全部按名字取值，不再有任何固定位置假设。
    """
    norm_to_canonical = {c.lower(): c for c in ALL_KNOWN_COLUMNS}
    col_map = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        key = str(h).strip().lower()
        canonical = norm_to_canonical.get(key)
        if canonical and canonical not in col_map:  # 同名列只认第一次出现的
            col_map[canonical] = i

    missing_required = [c for c in REQUIRED_COLUMNS if c not in col_map]
    if missing_required:
        raise SystemExit(
            f'表头里缺少必须的列: {missing_required}\n'
            f'表头原文: {[str(h) if h is not None else "" for h in header_row]}\n'
            f'本脚本只要求有 "nodeid" 这一列能被认出来，其余列（{OPTIONAL_COLUMNS}）都是\n'
            '可选的——缺哪个就当哪个信息没有，不需要凑齐全部列。如果你的表确实有\n'
            '对应的列只是名字不一样，改成上面列出的名字（忽略大小写/首尾空格）即可。'
        )
    print(f'识别到的列: {col_map}', file=sys.stderr)
    missing_optional = [c for c in OPTIONAL_COLUMNS if c not in col_map]
    if missing_optional:
        print(f'没识别到（按"没有这些信息"处理，不是报错）: {missing_optional}', file=sys.stderr)
    return col_map


def row_to_record(row, col_map):
    """把一行原始值按 col_map 转成 {列名: 值} 的字典，缺的列值为 None。"""
    rec = {}
    for name, idx in col_map.items():
        rec[name] = row[idx] if idx < len(row) else None
    return rec


def field(rec, name):
    """
    从 record 里取一个字段，统一转成字符串并去掉首尾空格，缺失/None 时返回
    空字符串。全脚本读取 结论/CANN算子/torch-api/不支持类型/失败大类 这些
    "语义上应该是文本"的字段都必须走这个函数，不能用 `(rec.get(x) or '')`
    这种写法——Excel 很容易把看起来像数字的单元格（比如有人只填了"8"想表示
    int8，没写成"int8"文本）自动存成 int/float 类型，`(8 or '').strip()`
    不会短路成空字符串（8 是真值），而是对整数调用 .strip()，直接崩溃。
    """
    v = rec.get(name)
    return '' if v is None else str(v).strip()


def _rows_from_xlsx(path, sheet_name):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise SystemExit(
            f'打不开 "{path}" 作为 Excel 文件：{e}\n'
            'openpyxl 只认 .xlsx/.xlsm/.xltx/.xltm，老的 .xls 二进制格式或损坏的文件都会在这里报错，'
            '不是脚本本身的问题，先确认文件能被 Excel 正常打开。'
        )
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
    except KeyError:
        raise SystemExit(f'sheet "{sheet_name}" 不存在。实际的 sheet 有: {wb.sheetnames}')
    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        raise SystemExit(f'sheet "{ws.title}" 是空的（连表头都没有），没法识别列，检查一下是不是传错了 sheet/文件。')
    col_map = build_column_map(header_row)
    raw_rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))
    return col_map, raw_rows


def _rows_from_csv(path):
    # SKILL.md 的"何时使用"明确写了"通常是 Excel/CSV 形式"，CSV 必须真的支持，
    # 不能只认 xlsx——用标准库 csv 模块读，不依赖 openpyxl（它只认 Excel 格式）。
    with open(path, encoding='utf-8-sig', newline='') as f:
        reader = list(csv.reader(f))
    if not reader:
        raise SystemExit(f'"{path}" 是空文件，连表头都没有。')
    col_map = build_column_map(reader[0])
    return col_map, reader[1:]


def load_records(input_path, sheet_name=None):
    ext = os.path.splitext(input_path)[1].lower()
    if ext == '.csv':
        col_map, raw_rows = _rows_from_csv(input_path)
    else:
        col_map, raw_rows = _rows_from_xlsx(input_path, sheet_name)

    records, malformed = [], []
    seen_nodeids = Counter()
    for i, r in enumerate(raw_rows, start=2):
        if not r or all(v is None or str(v).strip() == '' for v in r):
            continue  # 空行，静默跳过
        rec = row_to_record(r, col_map)
        nodeid_raw = rec.get('nodeid')
        if not nodeid_raw or '::' not in str(nodeid_raw):
            malformed.append((i, 'nodeid 列的值不像 nodeid（空或不含 "::"）', r))
            continue
        # 必须去首尾空格：复制粘贴很容易带进来，不去掉的话这个 nodeid 以后
        # 没法跟 pytest 实际收集到的 nodeid 精确匹配（用来配 skip-list 时会
        # 静默匹配失败），且同一个 nodeid 带不带空格会被后面的查重当成两条。
        rec['nodeid'] = str(nodeid_raw).strip()
        seen_nodeids[rec['nodeid']] += 1
        records.append(rec)
    if malformed:
        print(f'WARNING: {len(malformed)} 行格式异常，已跳过（不会让它们拖垮整个脚本），需人工复查：', file=sys.stderr)
        for row_num, reason, raw in malformed[:20]:
            print(f'  第 {row_num} 行: {reason} | {raw}', file=sys.stderr)
        if len(malformed) > 20:
            print(f'  ...还有 {len(malformed) - 20} 行未列出', file=sys.stderr)
    dups = {n: c for n, c in seen_nodeids.items() if c > 1}
    if dups:
        print(f'WARNING: {len(dups)} 个 nodeid 出现了不止一次（不会自动去重，原样全部保留），需确认是否符合预期：', file=sys.stderr)
        for n, c in list(dups.items())[:20]:
            print(f'  {n}  x{c}', file=sys.stderr)
        if len(dups) > 20:
            print(f'  ...还有 {len(dups) - 20} 个未列出', file=sys.stderr)
    if not records:
        print('WARNING: 有效数据行数为 0——多半是表头认出来了但内容格式不对，核实一下输入文件', file=sys.stderr)
    return records


def _md_cell(value):
    """
    把一个值安全地塞进 markdown 表格的一个单元格。两件事必须做：转义 "|"
    （否则被当成新的列分隔符，错位整行），把换行替换成空格（否则这一行会在
    渲染时断成两行，表格结构从这行往后全部错位）——换行是更隐蔽的那个：
    不报错、不崩溃，只是产出一张看起来正常、实则从某一行开始结构损坏的表，
    批量跑几千行时很容易被忽略过去。
    """
    s = str(value) if value is not None else ''
    s = s.replace('\r\n', ' ').replace('\r', ' ').replace('\n', ' ')
    return s.replace('|', '\\|')


def build_url_name_map(records):
    url_to_name = {}
    for rec in records:
        concl = field(rec, '结论')
        url = field(rec, '文档链接')
        m = NAME_PAT.match(concl)
        if m and url:
            url_to_name.setdefault(url, m.group(1))
    return url_to_name


def resolve_op_name(rec, url_to_name):
    concl = field(rec, '结论')
    m = NAME_PAT.match(concl)
    if m:
        return m.group(1), 'direct'
    url = field(rec, '文档链接')
    if url and url in url_to_name:
        return url_to_name[url], 'url_fallback'
    op_name = field(rec, 'CANN算子') or field(rec, 'torch-api')
    if op_name:
        return op_name, 'column_fallback'
    return None, 'unresolved'


def generate(input_path, out_md_path, sheet_name=None, memfmt_category='内存格式错误'):
    records = load_records(input_path, sheet_name)
    url_to_name = build_url_name_map(records)

    out_rows = []
    resolution_stats = Counter()
    unresolved = []

    memfmt_norm = (memfmt_category or '').strip()
    for rec in records:
        nodeid = rec.get('nodeid')
        op_name, how = resolve_op_name(rec, url_to_name)
        resolution_stats[how] += 1
        if op_name is None:
            unresolved.append(nodeid)
            op_name = '(未识别——需人工核实对应算子)'

        cat_norm = field(rec, '失败大类')
        dtype_val = field(rec, '不支持类型')
        unsupported = '非连续内存' if (memfmt_norm and cat_norm == memfmt_norm) else (dtype_val or '(未填写)')
        out_rows.append((nodeid, op_name, unsupported))

    lines = [
        '# 屏蔽用例汇总 三列表',
        '',
        f'> 数据来源：`{input_path}`',
        f'> 总计：{len(out_rows)} 条',
        '',
        '## 屏蔽汇总表',
        '',
        '| nodeid | 对应算子 | 不支持类型 |',
        '| --- | --- | --- |',
    ]
    for nodeid, op_name, unsupported in out_rows:
        lines.append(f'| {_md_cell(nodeid)} | {_md_cell(op_name)} | {_md_cell(unsupported)} |')

    with open(out_md_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))

    print(f'total rows: {len(out_rows)}', file=sys.stderr)
    print(f'resolution breakdown: {dict(resolution_stats)}', file=sys.stderr)
    if unresolved:
        print(f'WARNING: {len(unresolved)} rows could not resolve 对应算子, need manual review:', file=sys.stderr)
        for n in unresolved:
            print(f'  {n}', file=sys.stderr)
    print(f'written to {out_md_path}', file=sys.stderr)
    return len(out_rows), len(unresolved)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input_xlsx', help='已完成文档核验的中间表路径，.xlsx/.xlsm 或 .csv；只要求有 nodeid 列')
    parser.add_argument('output_md', help='输出的三列汇总 md 文件路径')
    parser.add_argument('--sheet', default=None, help='Sheet 名，默认用第一个激活的 sheet')
    parser.add_argument('--memfmt-category', default='内存格式错误', help='内存格式失败大类的名称，命中时"不支持类型"统一写"非连续内存"；传空字符串可关闭这条规则')
    args = parser.parse_args()

    total, unresolved_count = generate(args.input_xlsx, args.output_md, args.sheet, args.memfmt_category)
    if unresolved_count:
        sys.exit(1)
